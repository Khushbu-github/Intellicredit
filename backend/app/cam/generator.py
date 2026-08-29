"""
Credit Appraisal Memo generation.

Produces the document a credit committee actually reads: the recommendation
up front, then the working behind it - pillar-by-pillar scoring with the
evidence each component was scored on, the limit sizing, the risk register,
the covenant package and the processing history.
"""

import io
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

DASH = "—"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
CRITICAL = RGBColor(0xA1, 0x26, 0x26)
GOOD = RGBColor(0x14, 0x6C, 0x14)


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------

def _crore(value, decimals=2):
    if value is None:
        return DASH
    return f"Rs {value:,.{decimals}f} Cr"


def _times(value):
    return DASH if value is None else f"{value:.2f}x"


def _pct(value, decimals=1):
    return DASH if value is None else f"{value:.{decimals}f}%"


def _num(value, decimals=1):
    return DASH if value is None else f"{value:,.{decimals}f}"


def _text(value):
    return DASH if value in (None, "") else str(value)


def _when(value):
    if not value:
        return DASH
    if isinstance(value, str):
        return value
    return value.strftime("%d %b %Y, %H:%M")


# ---------------------------------------------------------------------------
# Document building blocks
# ---------------------------------------------------------------------------

def _configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(4)


def _heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = INK
    return heading


def _caption(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    run.italic = True
    return para


def _kv_table(doc, rows, widths=(2.6, 3.6)):
    """Two-column label/value table."""
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
    for i, width in enumerate(widths):
        for row in table.rows:
            row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def _grid_table(doc, headers, rows, widths=None):
    """Multi-column data table with a header row."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)

    if widths:
        for i, width in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Inches(width)

    doc.add_paragraph()
    return table


# ---------------------------------------------------------------------------
# Sections
# ---------------------------------------------------------------------------

def _cover(doc, data):
    title = doc.add_heading("Credit Appraisal Memo", level=0)
    for run in title.runs:
        run.font.color.rgb = INK

    sub = doc.add_paragraph()
    sub_run = sub.add_run("Confidential " + DASH + " for internal circulation only")
    sub_run.font.size = Pt(9)
    sub_run.font.color.rgb = MUTED
    sub_run.bold = True

    _kv_table(doc, [
        ("Borrower", _text(data.get("company_name"))),
        ("Case reference", f"IC-{data.get('id', 0):04d}"),
        ("Sector", _text(data.get("sector"))),
        ("Case opened", _when(data.get("created_at"))),
        ("Last analysed", _when(data.get("analysed_at"))),
        ("Prepared on", datetime.now().strftime("%d %b %Y, %H:%M")),
    ])


def _narrative_block(doc, narrative, key, heading):
    """Render one drafted section, if the model produced it."""
    body = (narrative or {}).get(key)
    if not body:
        return False
    _heading(doc, heading, 2)
    for paragraph in str(body).splitlines():
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())
    return True


def _executive_summary(doc, data):
    """
    The model-drafted opening, when one has been generated for this case.

    Omitted entirely rather than replaced with filler when no narrative
    exists, so a memo produced without model access simply starts at the
    recommendation instead of carrying an empty heading.
    """
    narrative = data.get("narrative")
    if not narrative:
        return

    _heading(doc, "Executive summary", 1)
    _narrative_block(doc, narrative, "executive_summary", "Overview")
    _narrative_block(doc, narrative, "risk_commentary", "Risk commentary")
    _narrative_block(doc, narrative, "mitigants", "Mitigants")

    _caption(
        doc,
        "Narrative drafted by "
        + str(narrative.get("generated_by", "a language model"))
        + " from the figures already determined by the policy engine. The "
        "recommendation, score, limit and covenants are computed "
        "deterministically and are not model output.",
    )


def _recommendation(doc, data):
    _heading(doc, "1. Recommendation", 1)

    decision = data.get("decision") or "PENDING"
    para = doc.add_paragraph()
    label = para.add_run(f"{decision}")
    label.bold = True
    label.font.size = Pt(14)
    label.font.color.rgb = CRITICAL if decision == "REJECT" else GOOD if decision == "APPROVE" else MUTED

    if data.get("recommended_limit_cr") is not None:
        amount = para.add_run(f"   {_crore(data['recommended_limit_cr'])}")
        amount.bold = True
        amount.font.size = Pt(14)

    doc.add_paragraph(_text(data.get("decision_reason")))

    _kv_table(doc, [
        ("Recommended limit", _crore(data.get("recommended_limit_cr"))),
        ("Pricing", _text(data.get("spread"))),
        ("Composite score", f"{_num(data.get('total_score'))} / 100"),
        ("Internal risk grade", _text(data.get("risk_grade"))),
        ("Existing exposure", _crore(data.get("existing_limits"))),
    ])

    gates = (data.get("raw_data") or {}).get("gates") or []
    if gates:
        _heading(doc, "1.1 Policy gates", 2)
        _grid_table(
            doc,
            ["Gate", "Result", "Assessment"],
            [
                [g["title"], "BREACH" if g["breached"] else "Pass", g["detail"]]
                for g in gates
            ],
            widths=(1.8, 0.9, 3.6),
        )

    _narrative_block(
        doc, data.get("narrative"), "recommendation_note", "1.2 Note to the sanctioning authority",
    )


def _financials(doc, data):
    _heading(doc, "2. Financial position", 1)

    revenue = data.get("revenue")
    ebitda = data.get("ebitda")
    margin = (ebitda / revenue * 100) if (revenue and ebitda) else None

    _kv_table(doc, [
        ("Revenue (annual)", _crore(revenue)),
        ("EBITDA", _crore(ebitda)),
        ("EBITDA margin", _pct(margin)),
        ("Debt service coverage", _times(data.get("dscr"))),
        ("Interest coverage", _times(data.get("interest_coverage"))),
        ("Debt / equity", _times(data.get("debt_equity"))),
        ("Net worth", _crore(data.get("net_worth"))),
        ("Promoter holding", _pct(data["promoter_stake"] * 100) if data.get("promoter_stake") is not None else DASH),
        ("Promoter pledge", _pct(data["promoter_pledge"] * 100) if data.get("promoter_pledge") is not None else DASH),
        ("Security cover", _times(data.get("security_cover"))),
        ("CIBIL commercial", _text(data.get("cibil_score"))),
        ("External rating", f"{_text(data.get('external_rating'))} / {_text(data.get('rating_outlook'))}"),
    ])
    _caption(doc, "A dash indicates the figure was not present in the documents on file.")

    benchmark = (data.get("research") or {}).get("benchmark")
    if benchmark and benchmark.get("comparisons"):
        _heading(doc, "2.1 Peer comparison", 2)
        _grid_table(
            doc,
            ["Metric", "Borrower", "Sector median", "Assessment"],
            [
                [
                    c["metric"],
                    f"{c['actual']}{c['unit']}" if c.get("available") else DASH,
                    f"{c['peer_median']}{c['unit']}",
                    c["verdict"],
                ]
                for c in benchmark["comparisons"]
            ],
            widths=(1.6, 1.0, 1.1, 2.6),
        )
        _caption(doc, benchmark.get("source", ""))


def _scoring(doc, data):
    _heading(doc, "3. 5Cs scoring", 1)

    pillars = [
        ("Character", data.get("character_score"), 25),
        ("Capacity", data.get("capacity_score"), 30),
        ("Capital", data.get("capital_score"), 20),
        ("Collateral", data.get("collateral_score"), 15),
        ("Conditions", data.get("conditions_score"), 10),
    ]

    _grid_table(
        doc,
        ["Pillar", "Score", "Weight", "Utilisation"],
        [
            [name, _num(score), str(weight),
             _pct((score / weight * 100) if score is not None else None, 0)]
            for name, score, weight in pillars
        ] + [["TOTAL", _num(data.get("total_score")), "100",
              _pct(data.get("total_score"), 0)]],
        widths=(1.8, 1.0, 1.0, 1.4),
    )

    breakdown = data.get("score_breakdown") or {}
    if breakdown:
        _heading(doc, "3.1 Scoring working", 2)
        for key, (name, _, _) in zip(
            ["character", "capacity", "capital", "collateral", "conditions"],
            pillars,
            strict=True,
        ):
            pillar = breakdown.get(key)
            if not pillar:
                continue

            para = doc.add_paragraph()
            run = para.add_run(f"{name} {DASH} {pillar['score']:.1f} of {pillar['weight']}")
            run.bold = True
            _caption(doc, pillar.get("intent", ""))

            _grid_table(
                doc,
                ["Component", "Points", "Max", "Evidence"],
                [
                    [
                        c["label"],
                        f"{c['points']:.1f}",
                        str(c["max"]),
                        c["basis"] + ("" if c["evidenced"] else "  [assumed - no evidence on file]"),
                    ]
                    for c in pillar.get("components", [])
                ],
                widths=(1.5, 0.7, 0.6, 3.5),
            )

        coverage = (data.get("raw_data") or {}).get("evidence_coverage")
        if coverage is not None:
            _caption(
                doc,
                f"Evidence coverage {coverage:.0f}%: the share of scoring components backed by "
                "a figure actually found in the case file. Components without evidence are "
                "scored at half their maximum and are marked above.",
            )


def _risk(doc, data):
    _heading(doc, "4. Risk register", 1)

    flags = data.get("red_flags") or []
    if flags:
        _grid_table(
            doc,
            ["Severity", "Finding", "Detail"],
            [[f["severity"].upper(), f["title"], f["detail"]] for f in flags],
            widths=(0.9, 1.8, 3.6),
        )
    else:
        doc.add_paragraph("No screening exceptions were raised against this case.")

    _heading(doc, "4.1 Screening measurements", 2)
    _kv_table(doc, [
        ("Circular trading overlap", _pct(data.get("gst_circular_trading_pct"))),
        ("Turnover mismatch (ITR vs GST)", _pct(data.get("gst_turnover_mismatch_pct"))),
        ("Litigation matters", _text(data.get("litigation_count"))),
        ("Dishonoured instruments", _text(data.get("bounce_count"))),
        ("Related-party dealings", "Yes" if data.get("related_party_flag") else "None noted"),
    ])

    research = data.get("research") or {}
    litigation = research.get("litigation") or {}
    if litigation.get("cases"):
        _heading(doc, "4.2 Litigation detail", 2)
        _grid_table(
            doc,
            ["Case", "Nature", "Jurisdiction", "Amount", "Status"],
            [
                [c["case_no"], c["nature"], c["jurisdiction"], _crore(c["amount_cr"]), c["status"]]
                for c in litigation["cases"]
            ],
            widths=(1.1, 1.9, 1.3, 0.9, 1.0),
        )
        _caption(doc, litigation.get("source", ""))

    mca = research.get("mca") or {}
    if mca.get("directors"):
        _heading(doc, "4.3 Directors and charges", 2)
        _grid_table(
            doc,
            ["Director", "DIN", "Appointed", "Status"],
            [
                [d["name"], d["din"], d["appointed"],
                 "DISQUALIFIED" if d["disqualified"] else "Active"]
                for d in mca["directors"]
            ],
            widths=(1.8, 1.2, 1.2, 1.3),
        )
        _kv_table(doc, [
            ("CIN", _text(mca.get("cin"))),
            ("Charges registered", _text(mca.get("charges_created"))),
            ("Charge value", _crore(mca.get("active_charges_value_cr"))),
            ("Annual filings", "In default" if mca.get("annual_filing_default") else "Up to date"),
        ])
        _caption(doc, mca.get("source", ""))


def _limit(doc, data):
    _heading(doc, "5. Limit assessment", 1)

    working = data.get("limit_working") or {}
    constraints = working.get("constraints") or []

    if constraints:
        _grid_table(
            doc,
            ["Method", "Basis", "Input", "Limit", "Binding"],
            [
                [
                    c["basis"],
                    c["formula"],
                    c["input"],
                    _crore(c["value_cr"]) if c.get("available") else "Not computable",
                    "YES" if c.get("binding") else "",
                ]
                for c in constraints
            ],
            widths=(1.2, 1.6, 1.9, 1.0, 0.7),
        )
        doc.add_paragraph(_text(working.get("note")))
    else:
        doc.add_paragraph("No limit sizing was performed.")

    pricing = (data.get("raw_data") or {}).get("pricing") or {}
    if pricing:
        _kv_table(doc, [
            ("Pricing band", _text(pricing.get("band"))),
            ("Applied spread", _text(pricing.get("spread"))),
            ("Basis", _text(pricing.get("basis"))),
        ])


def _covenants(doc, data):
    _heading(doc, "6. Covenants and conditions", 1)

    package = data.get("covenants") or {}

    financial = package.get("financial_covenants") or []
    if financial:
        _heading(doc, "6.1 Financial covenants", 2)
        _grid_table(
            doc,
            ["Covenant", "Threshold", "Rationale"],
            [[c["covenant"], c["threshold"], c["reason"]] for c in financial],
            widths=(1.7, 2.0, 2.6),
        )

    conditions = package.get("conditions_precedent") or []
    if conditions:
        _heading(doc, "6.2 Conditions precedent", 2)
        _grid_table(
            doc,
            ["Condition", "Rationale"],
            [[c["condition"], c["reason"]] for c in conditions],
            widths=(2.8, 3.5),
        )

    monitoring = package.get("monitoring") or []
    if monitoring:
        _heading(doc, "6.3 Ongoing monitoring", 2)
        _grid_table(
            doc,
            ["Requirement", "Frequency"],
            [[m["item"], m["frequency"]] for m in monitoring],
            widths=(4.2, 1.8),
        )


def _field_notes(doc, data):
    if data.get("factory_capacity") is None and not data.get("qualitative_notes"):
        return

    _heading(doc, "7. Relationship manager assessment", 1)
    _kv_table(doc, [
        ("Observed capacity utilisation",
         _pct(data["factory_capacity"] * 100, 0) if data.get("factory_capacity") is not None else DASH),
        ("Management quality",
         f"{data['mgmt_quality']} of 5" if data.get("mgmt_quality") else DASH),
    ])
    if data.get("qualitative_notes"):
        doc.add_paragraph(data["qualitative_notes"])


def _audit(doc, data):
    _heading(doc, "8. Documents and audit trail", 1)

    documents = data.get("documents") or []
    if documents:
        _heading(doc, "8.1 Documents on file", 2)
        _grid_table(
            doc,
            ["File", "Type", "Status", "Ingestion result"],
            [
                [d["filename"], d["doc_label"], d["parse_status"].upper(), d.get("parse_message") or ""]
                for d in documents
            ],
            widths=(1.7, 1.3, 0.9, 2.5),
        )

    events = data.get("events") or []
    if events:
        _heading(doc, "8.2 Processing history", 2)
        _grid_table(
            doc,
            ["Timestamp", "Trigger", "Outcome"],
            [[_when(e["at"]), e["trigger"], e["summary"]] for e in events],
            widths=(1.5, 1.2, 3.6),
        )


def _footer(doc):
    doc.add_paragraph()
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(
        "Generated by IntelliCredit from the documents on file. Scoring and decisioning are "
        "rules-based and deterministic. Registry, litigation, news and peer-benchmark inputs "
        "are simulated in this build and are not verified third-party data. This memo is a "
        "decision aid and does not constitute a sanction."
    )
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    run.italic = True


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def generate_cam_docx(app_data):
    """Build the memo and return it as an in-memory .docx stream."""
    doc = Document()
    _configure_styles(doc)

    _cover(doc, app_data)
    _executive_summary(doc, app_data)
    _recommendation(doc, app_data)
    _financials(doc, app_data)
    _scoring(doc, app_data)
    _risk(doc, app_data)
    _limit(doc, app_data)
    _covenants(doc, app_data)
    _field_notes(doc, app_data)
    _audit(doc, app_data)
    _footer(doc)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
