from docx import Document
from datetime import datetime
import io

def generate_cam_docx(app_data):
    doc = Document()
    
    doc.add_heading('Credit Appraisal Memo (CAM)', 0)
    
    # 1. Company Profile
    doc.add_heading('1. Company Profile', level=1)
    doc.add_paragraph(f"Company ID: {app_data.get('id')}")
    doc.add_paragraph(f"Status: {app_data.get('status')}")
    
    # 2. Financial Summary
    doc.add_heading('2. Financial Summary', level=1)
    doc.add_paragraph(f"Revenue: ₹{app_data.get('revenue', 0)} Cr")
    doc.add_paragraph(f"EBITDA: ₹{app_data.get('ebitda', 0)} Cr")
    doc.add_paragraph(f"DSCR: {app_data.get('dscr', 0)}x")
    
    # 3. 5Cs Scoring Analysis
    doc.add_heading('3. 5Cs Scoring Analysis', level=1)
    doc.add_paragraph(f"Character Score: {app_data.get('character_score', 0)}/25")
    doc.add_paragraph(f"Capacity Score: {app_data.get('capacity_score', 0)}/30")
    doc.add_paragraph(f"Capital Score: {app_data.get('capital_score', 0)}/20")
    doc.add_paragraph(f"Collateral Score: {app_data.get('collateral_score', 0)}/15")
    doc.add_paragraph(f"Conditions Score: {app_data.get('conditions_score', 0)}/10")
    doc.add_paragraph(f"TOTAL SCORE: {app_data.get('total_score', 0)}/100")
    
    # 4. GST & Bank Analysis
    doc.add_heading('4. GST & Bank Analysis', level=1)
    doc.add_paragraph(f"Circular Trading Pct: {app_data.get('gst_circular_trading_pct', 0)}%")
    
    # 5. Final Decision
    doc.add_heading('5. Final Decision', level=1)
    doc.add_paragraph(f"Final Decision: {app_data.get('decision', 'N/A')}")
    doc.add_paragraph(f"Reasoning: {app_data.get('decision_reason', 'N/A')}")
    doc.add_paragraph(f"Recommended Limit: ₹{app_data.get('recommended_limit_cr', 0)} Cr")
    doc.add_paragraph(f"Applied Spread: {app_data.get('spread', 'N/A')}")
    
    doc.add_paragraph(f"\nReport Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    # Use a buffer instead of a file
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
